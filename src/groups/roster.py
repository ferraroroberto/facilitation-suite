"""The roster and the breakout rounds of one session (epic §13).

- **Roster**: ``roster.xlsx`` in the session folder (imported from anywhere:
  the original is copied, never edited). Columns by header, any order:
  ``name`` (required), ``role``, ``company``, ``country``, ``present``,
  ``email`` (optional). Without a ``name`` header the first column is the
  name. The file's ``present`` is only the starting point: the Groups tab's
  switches are authoritative and saved in ``groups.yaml``.
- **Rounds**: ``groups.yaml`` — who was absent, the three rounds from
  ``build_groups`` and when they were shuffled. Plain YAML, hand-editable.
- **Zoom**: the copy-paste text for manual room assignment, and the
  pre-assign CSV (``Pre-assign Room Name,Email Address``) when every present
  participant has an email — otherwise the list of who is missing one.
"""

from __future__ import annotations

import csv
import io
import logging
import random
import shutil
import unicodedata
from dataclasses import asdict, dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

import yaml
from openpyxl import load_workbook

from src.groups.shuffle import _assign_ids, build_groups
from src.sessions.store import atomic_write_text

logger = logging.getLogger(__name__)

ROSTER_FILE = "roster.xlsx"
GROUPS_FILE = "groups.yaml"
ROUNDS = ("pairs", "g4a", "g4b")
ROUND_LABELS = {"pairs": "Pairs", "g4a": "Groups of 4 · A", "g4b": "Groups of 4 · B"}
COLUMNS = ("name", "role", "company", "country", "present", "email")


class RosterError(Exception):
    def __init__(self, status: int, code: str, message: str) -> None:
        super().__init__(message)
        self.status = status
        self.code = code


@dataclass
class Person:
    name: str
    role: str = ""
    company: str = ""
    country: str = ""
    email: str = ""
    present: bool = True


def sort_key(name: str) -> str:
    """Alphabetical with accents ignored (Á sorts with A)."""
    return "".join(c for c in unicodedata.normalize("NFD", name.lower()) if unicodedata.category(c) != "Mn")


def _present(value: object) -> bool:
    if value is None or value == "":
        return True  # no column / empty cell: assume present, the switches decide
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return int(value) == 1
    return str(value).strip().lower() in ("1", "yes", "y", "si", "sí", "true", "x")


def read_roster(path: Path) -> list[Person]:
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — openpyxl raises several types on a bad file
        raise RosterError(422, "roster_unreadable", f"The roster could not be read: {exc}") from exc
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    headers = [str(h or "").strip().lower() for h in rows[0]]
    col = {c: headers.index(c) for c in COLUMNS if c in headers}
    name_col = col.get("name", 0)
    people: dict[str, Person] = {}
    for row in rows[1:]:
        cell = row[name_col] if name_col < len(row) else None
        name = str(cell).strip() if cell is not None else ""
        if not name:
            continue

        def get(key: str, _row: tuple = row) -> str:
            i = col.get(key)
            v = _row[i] if i is not None and i < len(_row) else None
            return str(v).strip() if v is not None else ""

        present_cell = row[col["present"]] if "present" in col and col["present"] < len(row) else None
        people[name] = Person(name, get("role"), get("company"), get("country"), get("email"), _present(present_cell))
    return sorted(people.values(), key=lambda p: sort_key(p.name))


def import_roster(folder: Path, source: Path) -> list[Person]:
    """Copy an xlsx into the session folder as roster.xlsx (validated first)."""
    if not source.is_file():
        raise RosterError(404, "not_found", "That file does not exist")
    if source.suffix.lower() != ".xlsx":
        raise RosterError(422, "not_xlsx", "The roster must be an .xlsx file")
    people = read_roster(source)
    if not people:
        raise RosterError(422, "roster_empty", "No names found in the first sheet")
    target = folder / ROSTER_FILE
    if source.resolve() != target.resolve():
        tmp = target.with_suffix(".xlsx.tmp")
        shutil.copyfile(source, tmp)
        tmp.replace(target)
    # A new roster starts from its own present column: forget earlier switches.
    state = load_groups(folder)
    state["absent"] = [p.name for p in people if not p.present]
    save_groups(folder, state)
    logger.info("ℹ️ roster imported: %d people (%d absent)", len(people), len(state["absent"]))
    return people


def load_groups(folder: Path) -> dict[str, Any]:
    path = folder / GROUPS_FILE
    if not path.is_file():
        return {"absent": [], "rounds": None, "shuffled_at": None}
    try:
        data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    except (OSError, yaml.YAMLError) as exc:
        raise RosterError(422, "groups_unreadable", f"groups.yaml could not be read: {exc}") from exc
    data.setdefault("absent", [])
    data.setdefault("rounds", None)
    data.setdefault("shuffled_at", None)
    return data


def save_groups(folder: Path, data: dict[str, Any]) -> None:
    atomic_write_text(folder / GROUPS_FILE, yaml.safe_dump(data, sort_keys=False, allow_unicode=True, width=100))


def roster_state(folder: Path) -> tuple[list[Person], dict[str, Any]]:
    """The roster with presence from groups.yaml (the switches win over the file)."""
    path = folder / ROSTER_FILE
    people = read_roster(path) if path.is_file() else []
    state = load_groups(folder)
    if not (folder / GROUPS_FILE).is_file():
        # nothing toggled yet: the file's own present column is the starting point
        state["absent"] = [p.name for p in people if not p.present]
    absent = set(state["absent"])
    for p in people:
        p.present = p.name not in absent
    return people, state


def set_present(folder: Path, name: str, present: bool) -> None:
    people, state = roster_state(folder)
    if name not in {p.name for p in people}:
        raise RosterError(404, "unknown_person", f"{name!r} is not on the roster")
    absent = [n for n in state["absent"] if n != name] + ([] if present else [name])
    state["absent"] = sorted(absent, key=sort_key)
    save_groups(folder, state)


def shuffle(folder: Path, seed: Optional[int] = None) -> dict[str, Any]:
    people, state = roster_state(folder)
    present = [p.name for p in people if p.present]
    if len(present) < 2:
        raise RosterError(422, "too_few", "At least two present participants are needed")
    if seed is not None:
        random.seed(seed)
    rounds = build_groups(present)
    state["rounds"] = rounds
    state["shuffled_at"] = datetime.now().isoformat(timespec="seconds")
    state["present_count"] = len(present)
    save_groups(folder, state)
    logger.info("✅ groups shuffled: %d present → %s", len(present), {k: len(v) for k, v in rounds.items()})
    return state


def mixing(rounds: dict[str, list[list[str]]]) -> int:
    """How many people share a round-B room with someone from their round-A room."""
    a = {p: set(g) for g in rounds["g4a"] for p in g}
    return sum(1 for g in rounds["g4b"] for p in g if (set(g) & a.get(p, set())) - {p})


def zoom_text(groups: list[list[str]], label: str = "Room") -> str:
    """Plain text to paste while assigning Zoom breakout rooms by hand."""
    return "\n".join(f"{label} {i}: {', '.join(g)}" for i, g in enumerate(groups, 1))


def zoom_csv(people: list[Person], groups: list[list[str]]) -> tuple[Optional[str], list[str]]:
    """Zoom's pre-assign CSV, or (None, names without an email)."""
    email = {p.name: p.email for p in people}
    names = [n for g in groups for n in g]
    missing = [n for n in names if not email.get(n)]
    if missing:
        return None, missing
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\r\n")
    w.writerow(["Pre-assign Room Name", "Email Address"])
    for room, n in sorted(((room, n) for n, room in _assign_ids(groups).items()), key=lambda x: (x[0], sort_key(x[1]))):
        w.writerow([f"Room {room}", email[n]])
    return buf.getvalue(), []


def payload(folder: Path) -> dict[str, Any]:
    """Everything the Groups tab shows."""
    people, state = roster_state(folder)
    rounds = state.get("rounds")
    out: dict[str, Any] = {
        "roster": [asdict(p) for p in people],
        "present": sum(1 for p in people if p.present),
        "total": len(people),
        "has_roster": (folder / ROSTER_FILE).is_file(),
        "shuffled_at": state.get("shuffled_at"),
        "rounds": rounds,
        "labels": ROUND_LABELS,
    }
    if rounds:
        present_now = {p.name for p in people if p.present}
        in_rounds = {n for g in rounds["pairs"] for n in g}
        out["stale"] = present_now != in_rounds  # presence changed since the shuffle
        out["mixing"] = mixing(rounds)
        out["texts"] = {r: zoom_text(rounds[r]) for r in ROUNDS}
        out["missing_emails"] = zoom_csv(people, rounds["pairs"])[1]
    return out
