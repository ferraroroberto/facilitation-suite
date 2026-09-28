"""Import Kahoot's spreadsheet template (``KahootQuizTemplate.xlsx``) into quiz plan items.

The official template has one sheet: a logo, the text ``Quiz template`` and
instructions in rows 1–7, the header in row 8 (columns B–H: ``Question - max
120 characters``, ``Answer 1 - max 75 characters`` … ``Answer 4``, ``Time
limit (sec) – 5, 10, …``, ``Correct answer(s) - choose at least one``) and up
to 100 questions from row 9. Column A numbers the rows even when they are
empty, so the questions end at the first empty question cell.

The header is found by its texts' prefixes in the first ``HEADER_ROWS`` rows,
not at a fixed cell. ``Correct answer(s)`` is a number (``1``) or a comma list
(``"1,3"``). The template has no quiz name: the caller passes a title.

Only this layout is in scope; Kahoot's report ``.xlsx`` is not (#50).
"""

from __future__ import annotations

import logging
import math
from pathlib import Path
from typing import Any, Optional

from openpyxl import load_workbook

from src.quiz.model import ANSWERS, InvalidQuestion, QuizError, QuizQuestion, make_question
from src.sessions.model import Item, Section, Session, ensure_ids, new_id

logger = logging.getLogger(__name__)

HEADER_ROWS = 20
# column key → the header text's start (lower case)
HEADERS: dict[str, str] = {
    "question": "question",
    **{f"answer_{i + 1}": f"answer {i + 1}" for i in range(ANSWERS)},
    "time_limit": "time limit",
    "correct": "correct answer",
}
REQUIRED = ("question", "answer_1", "answer_2", "correct")
QUIZ_PROFILE = "camera_pip"  # as a new activity in the Plan tab
OVERHEAD_S = 30  # per question, besides its time limit: the reveal and the leaderboard


def _header(row: tuple[Any, ...]) -> dict[str, int]:
    """Column index of each known header in ``row`` (empty when it is not the header row)."""
    found: dict[str, int] = {}
    for i, cell in enumerate(row):
        text = str(cell).strip().lower() if isinstance(cell, str) else ""
        for key, prefix in HEADERS.items():
            if key not in found and text.startswith(prefix):
                found[key] = i
    return found if all(k in found for k in REQUIRED) else {}


def read_kahoot_xlsx(path: Path) -> list[QuizQuestion]:
    """Every question of a Kahoot template, checked; raises ``QuizError`` naming the first bad row."""
    if not path.is_file():
        raise QuizError(404, "not_found", "That file does not exist")
    if path.suffix.lower() != ".xlsx":
        raise QuizError(422, "not_xlsx", "The Kahoot quiz must be an .xlsx file")
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — openpyxl raises several types on a bad file
        raise QuizError(422, "quiz_unreadable", f"The quiz file could not be read: {exc}") from exc
    try:
        rows = wb.worksheets[0].iter_rows(min_row=1, values_only=True)
        cols: dict[str, int] = {}
        questions: list[QuizQuestion] = []
        for number, row in enumerate(rows, start=1):
            if not cols:
                cols = _header(row)
                if not cols and number >= HEADER_ROWS:
                    break
                continue

            values = {key: row[i] if i < len(row) else None for key, i in cols.items()}
            if not str(values["question"] or "").strip():
                break  # the questions end at the first empty question cell
            try:
                questions.append(make_question(values["question"],
                                               [values.get(f"answer_{i + 1}") for i in range(ANSWERS)],
                                               values["correct"], values.get("time_limit")))
            except InvalidQuestion as exc:
                raise QuizError(422, "quiz_row_invalid", f"Row {number}: {exc}",
                                {"row": number, "reason": exc.reason}) from exc
    finally:
        wb.close()
    if not cols:
        raise QuizError(422, "quiz_no_header",
                        f"No Kahoot template header (Question, Answer 1, … Correct answer(s)) in the first {HEADER_ROWS} rows")
    if not questions:
        raise QuizError(422, "quiz_empty", "The quiz has no questions")
    logger.info("ℹ️ Kahoot template %s: %d questions", path.name, len(questions))
    return questions


def build_items(title: str, questions: list[QuizQuestion]) -> list[Item]:
    """A game in plan order: the lobby, one ``quiz`` item per question, the podium."""
    def activity(type_: str, **fields: Any) -> Item:
        return Item(kind="activity", id=new_id("act"), type=type_, profile=QUIZ_PROFILE, **fields)

    return [
        activity("quiz_lobby", options={"title": title}),
        *(activity("quiz", question=q.question, options=q.options()) for q in questions),
        activity("quiz_podium"),
    ]


def planned_minutes(questions: list[QuizQuestion]) -> int:
    """A first estimate for the section: every time limit plus the reveal and leaderboard after it."""
    return max(1, math.ceil(sum(q.time_limit + OVERHEAD_S for q in questions) / 60))


def import_kahoot(session: Session, path: Path, title: str, section_id: Optional[str] = None) -> dict[str, Any]:
    """Read ``path`` and add its game to ``session``: to the end of ``section_id``, else as a new
    section named ``title`` at the end of the plan. Nothing changes when the file is refused."""
    questions = read_kahoot_xlsx(path)
    items = build_items(title, questions)
    if section_id:
        section = next((s for s in session.sections if s.id == section_id), None)
        if section is None:
            raise QuizError(404, "section_not_found", "No such section in this plan")
        section.items.extend(items)
    else:
        section = Section(name=title, minutes=planned_minutes(questions), items=items)
        session.sections.append(section)
    ensure_ids(session)
    logger.info("✅ Kahoot quiz %r imported: %d questions into section %r", title, len(questions), section.name)
    return {"section_id": section.id, "section": section.name, "questions": len(questions),
            "items": len(items), "lobby_id": items[0].id}
