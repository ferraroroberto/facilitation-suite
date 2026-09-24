"""The Excel report (epic §14): a summary, one sheet per activity, participation.

- **Summary** — the session, then one row per activity: answers, people,
  hidden, when it was captured.
- **One sheet per activity** — every answer with the person's name, the chat
  time, the text, its parsed value, and whether it was hidden.
- **Participation** — per person: answers counted in each activity, the
  total, and chat messages sent (the StreamAlive "fans" idea), most active
  first.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
BOLD = Font(bold=True)


def sheet_title(n: int, title: str, taken: set[str]) -> str:
    base = BAD_SHEET_CHARS.sub(" ", f"{n} {title}").strip()[:31].rstrip() or str(n)
    name, k = base, 2
    while name.casefold() in taken or name.casefold() in ("summary", "participation"):
        suffix = f" ({k})"
        name, k = base[: 31 - len(suffix)] + suffix, k + 1
    taken.add(name.casefold())
    return name


def _table(ws, header: list[str], rows: list[list[Any]], widths: list[int]) -> None:  # noqa: ANN001 — openpyxl sheet
    ws.append(header)
    for c in ws[ws.max_row]:
        c.font = BOLD
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_report(results: dict[str, Any], people: list[dict[str, Any]], out: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    sess = results["session"]
    ws.append([sess["title"]])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([(sess.get("date") or "")[:10]])
    ws.append([])
    _table(ws, ["#", "Activity", "Type", "Answers", "People", "Hidden", "Captured", "Result"],
           [[a["number"], a["title"], a["type_label"], a["answers"], a["people"], a["hidden"], a["captured"], a["summary"]]
            for a in results["activities"]],
           [5, 48, 14, 10, 10, 10, 14, 36])
    ws.freeze_panes = "A5"

    taken: set[str] = set()
    for a in results["activities"]:
        sh = wb.create_sheet(sheet_title(a["number"], a["title"], taken))
        sh.append([a["title"]])
        sh["A1"].font = Font(bold=True, size=13)
        if a["question"] and a["question"] != a["title"]:
            sh.append([a["question"]])
        sh.append([f"{a['answers']} answers from {a['people']} people · captured {a['captured']}"
                   + (f" · {a['hidden']} hidden" if a["hidden"] else "")])
        sh.append([])
        first = sh.max_row + 1
        _table(sh, ["Name", "Time", "Answer", "Parsed", "Hidden"],
               [[r["sender"], r["time"], r["text"], r["value"], "yes" if r["hidden"] else ""] for r in a["answer_rows"]],
               [26, 8, 60, 30, 8])
        sh.freeze_panes = f"A{first + 1}"

    ps = wb.create_sheet("Participation")
    acts = results["activities"]
    _table(ps, ["Name", "Answers", *[f"{a['number']} {a['title']}"[:40] for a in acts], "Chat messages"],
           [[p["name"], p["total"], *[p["by_activity"].get(a["id"], 0) for a in acts], p["messages"]] for p in people],
           [26, 10, *[14 for _ in acts], 14])
    ps.freeze_panes = "B2"

    out.parent.mkdir(parents=True, exist_ok=True)
    tmp = out.with_suffix(".xlsx.tmp")
    wb.save(tmp)
    tmp.replace(out)
    return out
