"""Results (step 12): the PDF order, the Excel report, the Zoom-chat check."""

from __future__ import annotations

import json
import logging
import re
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from src.activities.registry import options_with_defaults, parser, result_for
from src.geo.gazetteer import gazetteer
from src.results.aliases import AliasError, Aliases, read_aliases, suggest, write_aliases
from src.results.collect import load_results, participation, participation_table, timeline
from src.results.excel import sheet_title
from src.results.pdf import activities_html, activities_pages, session_html, share_html, share_pages
from src.results.reconcile import minutes_of, norm_text, parse_saved_chat, reconcile
from src.sessions.model import parse_session
from tests.fixtures.demo import LONG_TOP_ANSWER, PLAN, build_demo_session, write_demo_run, zoom_saved_chat


def _run(isolated_env: Path) -> tuple[str, Path, dict]:
    sid, folder = build_demo_session(isolated_env / "sessions" / "demo" / "results", isolated_env / "sessions.local.yaml")
    return sid, folder, write_demo_run(folder)


# ---- the order of the session PDF ---------------------------------------------

ITEMS = [
    {"id": "slide-1", "kind": "slide", "title": "One", "slide_file": "slide-1.png"},
    {"id": "act-a", "kind": "activity", "title": "A"},
    {"id": "slide-2", "kind": "slide", "title": "Two", "slide_file": "slide-2.png"},
    {"id": "brk", "kind": "break", "title": "Break"},
]


def ev(name: str, item: str | None = None, **kw) -> dict:
    return {"event": name, "at": "t", **({"item_id": item} if item else {}), **kw}


def test_slides_at_first_showing_and_captures_at_their_last_run() -> None:
    events = [ev("session_live", "slide-1"), ev("item", "act-a"), ev("capture_start", "act-a"), ev("capture_stop", "act-a"),
              ev("item", "slide-2"), ev("item", "act-a"), ev("capture_reopen", "act-a"), ev("capture_stop", "act-a"),
              ev("item", "slide-1"), ev("item", "brk"), ev("item", "slide-gone")]
    pages = timeline(events, ITEMS, {"act-a"})
    assert [(p["kind"], p["item_id"]) for p in pages] == [("slide", "slide-1"), ("slide", "slide-2"), ("capture", "act-a")]


def test_a_capture_sits_where_it_was_run_not_where_the_next_item_stopped_it() -> None:
    # A capture stops when the facilitator leaves (the stop is logged after the next item), so the stop
    # lands behind the slide shown next; the page belongs at the run, planned order notwithstanding.
    items = [{"id": f"slide-{n}", "kind": "slide", "title": str(n), "slide_file": f"slide-{n}.png"} for n in (1, 2, 3)]
    items += [{"id": "act-a", "kind": "activity", "title": "A"}, {"id": "act-b", "kind": "activity", "title": "B"}]
    events = [ev("session_live", "slide-3"),
              ev("item", "act-b"), ev("capture_start", "act-b"), ev("item", "slide-1"), ev("capture_stop", "act-b"),
              ev("item", "act-a"), ev("capture_start", "act-a"), ev("item", "slide-2"), ev("capture_stop", "act-a")]
    pages = timeline(events, items, {"act-a", "act-b"})
    assert [p["item_id"] for p in pages] == ["slide-3", "act-b", "slide-1", "act-a", "slide-2"]


def test_a_reopened_capture_moves_to_where_it_was_reopened() -> None:
    events = [ev("session_live", "slide-1"), ev("item", "act-a"), ev("capture_start", "act-a"), ev("item", "slide-2"),
              ev("capture_stop", "act-a"), ev("item", "act-a"), ev("capture_reopen", "act-a"), ev("capture_stop", "act-a")]
    pages = timeline(events, ITEMS, {"act-a"})
    assert [p["item_id"] for p in pages] == ["slide-1", "slide-2", "act-a"]


def test_an_old_log_without_the_starting_item_starts_on_item_one() -> None:
    pages = timeline([ev("session_live"), ev("item", "slide-2"), ev("session_live")], ITEMS, set())
    assert [p["item_id"] for p in pages] == ["slide-1", "slide-2"]


def test_a_capture_without_its_frozen_file_is_not_a_page() -> None:
    assert timeline([ev("capture_stop", "act-a")], ITEMS, set()) == []


# ---- the results of a finished run ---------------------------------------------

def test_the_pdf_follows_the_run_when_each_stop_is_logged_after_the_next_item(isolated_env: Path) -> None:
    _, folder, run = _run(isolated_env)
    events = list(run["events"])
    i = 0
    while i < len(events) - 1:  # stop-on-leave writes a capture_stop just after the item that ended it
        if events[i]["event"] == "capture_stop" and events[i + 1]["event"] == "item":
            events[i], events[i + 1] = events[i + 1], events[i]
            i += 1
        i += 1
    assert events != run["events"]
    (folder / "live" / "events.jsonl").write_text("".join(json.dumps(e) + "\n" for e in events), encoding="utf-8")
    frozen = {p.stem for p in (folder / "live" / "captures").glob("*.json")}
    ran = [e["item_id"] for e in run["events"] if e["event"] in ("session_live", "item") and (e["item_id"].startswith("slide") or e["item_id"] in frozen)]
    pages = load_results(folder, parse_session(PLAN))["pages"]
    assert [p["item_id"] for p in pages if p["kind"] in ("slide", "capture")] == ran


def test_results_read_back_in_the_order_they_happened(isolated_env: Path) -> None:
    _, folder, run = _run(isolated_env)
    data = load_results(folder, parse_session(PLAN))
    assert [a["id"] for a in data["activities"]] == ["act-map", "act-weather", "act-kryptonite", "act-enemy", "act-ideas", "act-takeaway"]
    assert (data["slides"], data["captures"]) == (10, 6)
    assert data["pages"][0] == {**data["pages"][0], "kind": "slide", "item_id": "slide-101"}
    assert data["pages"][1]["kind"] == "capture" and data["pages"][1]["item_id"] == "act-map"

    kr = next(a for a in data["activities"] if a["id"] == "act-kryptonite")
    assert (kr["answers"], kr["hidden"]) == (8, 1)
    assert kr["summary"] == "8 answers · word cloud" and kr["top_label"] == "Top words"
    # #99: kryptonite is verbatim — its top (and tile) answer is the 60-char phrase, not "meetings"
    assert kr["top"][0] == {"label": LONG_TOP_ANSWER, "count": 3}
    assert kr["tile"] == LONG_TOP_ANSWER
    assert next(r for r in kr["answer_rows"] if r["text"] == "hello")["hidden"] is True
    weather = next(a for a in data["activities"] if a["id"] == "act-weather")
    assert weather["summary"].startswith("8 answers · average 3.") and weather["tile"].startswith("3.")
    geo = next(a for a in data["activities"] if a["id"] == "act-map")
    assert "1 unplaced" in geo["summary"] and geo["top_label"] == "Countries"
    assert next(r for r in geo["answer_rows"] if r["text"] == "Grnada")["value"] == "unplaced"

    people = participation(data["activities"], run["chat"])
    assert sum(p["total"] for p in people) == 39  # 40 answers, one hidden
    assert "You" not in {p["name"] for p in people}
    assert people[0]["total"] >= people[-1]["total"]


def _map_rows(folder: Path, rows: list[tuple[str, str]], **manual: str) -> dict[str, str]:
    """The report's value per answer text, for a map capture rebuilt from ``rows`` (sender, text); ``manual`` maps text → geonameid."""
    path = folder / "live" / "captures" / "act-map.json"
    frozen = json.loads(path.read_text(encoding="utf-8"))
    msgs = [{"id": i + 1, "sender": s, "text": t, "time": "", "received_at": 0, **({"place_id": manual[t]} if t in manual else {})}
            for i, (s, t) in enumerate(rows)]
    opts = options_with_defaults("map", frozen["item"].get("options") or {})
    frozen["answers"] = [{**m, "hidden": False, "parsed": parser("map").parse(m, opts)} for m in msgs]
    frozen["result"] = result_for("map", frozen["item"].get("options") or {}, msgs)
    path.write_text(json.dumps(frozen), encoding="utf-8")
    geo = next(a for a in load_results(folder, parse_session(PLAN))["activities"] if a["id"] == "act-map")
    return {r["text"]: r["value"] for r in geo["answer_rows"]}


def test_the_report_places_an_answer_where_the_room_put_it_on_the_map(isolated_env: Path) -> None:
    _, folder, _ = _run(isolated_env)
    room = [("Ana", "Madrid, España"), ("Sam", "Sevilla, España"), ("Leo", "Valencia")]
    assert _map_rows(folder, room)["Valencia"] == "Valencia, España"
    assert _map_rows(folder, room[-1:])["Valencia"] == "Valencia, Venezuela"  # alone, the larger one


def test_a_hand_fixed_or_unplaced_answer_keeps_its_report_value(isolated_env: Path) -> None:
    _, folder, _ = _run(isolated_env)
    caracas = next(p for p in gazetteer().suggest("Caracas") if p.country == "VE")
    rows = _map_rows(folder, [("Ana", "Madrid, España"), ("Sam", "Sevilla, España"), ("Leo", "Valencia"), ("Bea", "zzzz")],
                     Valencia=caracas.geonameid)
    assert (rows["Valencia"], rows["zzzz"]) == ("Caracas, Venezuela", "unplaced")


def test_a_session_that_never_went_live_has_no_results(isolated_env: Path) -> None:
    _, folder = build_demo_session(isolated_env / "sessions" / "demo" / "fresh")
    data = load_results(folder, parse_session(PLAN))
    assert data == {**data, "activities": [], "pages": [], "went_live": False}


def test_a_damaged_capture_is_left_out_not_fatal(isolated_env: Path) -> None:
    _, folder, _ = _run(isolated_env)
    (folder / "live" / "captures" / "act-enemy.json").write_text("{not json", encoding="utf-8")
    with open(folder / "live" / "events.jsonl", "a", encoding="utf-8") as fh:
        fh.write("garbage\n")
    data = load_results(folder, parse_session(PLAN))
    assert "act-enemy" not in [a["id"] for a in data["activities"]] and len(data["activities"]) == 5


# ---- API: results, frozen image, Excel -------------------------------------------

def test_results_api_and_the_frozen_image(client, isolated_env: Path) -> None:
    sid, _, _ = _run(isolated_env)
    data = client.get(f"/api/sessions/{sid}/results").json()
    assert data["went_live"] and len(data["activities"]) == 6
    assert data["reconciliation"] is None and data["exports"] == {"pdf": None, "share": None, "activities": None, "xlsx": None}
    r = client.get(f"/api/sessions/{sid}/results/captures/act-kryptonite.png")
    assert r.status_code == 200 and r.content[:4] == b"\x89PNG"
    assert client.get(f"/api/sessions/{sid}/results/captures/nope.png").status_code == 404
    assert client.get(f"/api/sessions/{sid}/results/captures/..%2Fx.png").status_code == 404


def test_excel_report_has_a_sheet_per_activity_and_participation(client, isolated_env: Path) -> None:
    sid, folder, _ = _run(isolated_env)
    assert client.get(f"/api/sessions/{sid}/exports/report.xlsx").status_code == 404  # nothing built yet
    r = client.post(f"/api/sessions/{sid}/exports/xlsx")
    assert r.status_code == 200 and r.json()["file"] == "report.xlsx"
    opened = client.get(f"/api/sessions/{sid}/exports/report.xlsx")
    assert opened.status_code == 200 and opened.content[:2] == b"PK"
    assert opened.headers["content-disposition"].startswith("inline")  # opens in a tab, never lands in Downloads
    wb = load_workbook(folder / "exports" / "report.xlsx")
    assert wb.sheetnames[0] == "Summary" and wb.sheetnames[-1] == "Participation"
    assert len(wb.sheetnames) == 8 and all(len(n) <= 31 for n in wb.sheetnames)
    kr = wb[next(n for n in wb.sheetnames if n.startswith("3 "))]
    rows = [[c.value for c in row] for row in kr.iter_rows()]
    header = rows.index(["Name", "Time", "Answer", "Parsed", "Hidden"])
    assert ["yes"] == [r[4] for r in rows[header + 1:] if r[2] == "hello"]
    part = [[c.value for c in row] for row in wb["Participation"].iter_rows()]
    assert part[0][0] == "Name" and part[0][-1] == "Chat messages" and len(part[0]) == 11
    assert part[0][-3:-1] == ["Early (1–3)", "Late (4–6)"]
    assert all(r[-3] + r[-2] == r[1] for r in part[1:])  # every counted answer is early or late
    for name in ("session.pdf", "session-full.pdf", "session-share.pdf", "session-activities.pdf"):
        assert client.get(f"/api/sessions/{sid}/exports/{name}").status_code == 404


def test_excel_needs_a_capture(client, isolated_env: Path) -> None:
    sid, _ = build_demo_session(isolated_env / "sessions" / "demo" / "empty", isolated_env / "sessions.local.yaml")
    assert client.post(f"/api/sessions/{sid}/exports/xlsx").json()["error"]["code"] == "nothing_to_export"
    assert client.post(f"/api/sessions/{sid}/exports/pdf").json()["error"]["code"] == "nothing_to_export"


# ---- the participation table: early vs late, one helper for the PDF and the Excel sheet ----

def _answers(*senders: str) -> list[dict]:
    return [{"sender": s, "hidden": False} for s in senders]


def _acts(per_activity: list[list[str]]) -> list[dict]:
    return [{"id": f"a{n}", "number": n, "title": f"Activity {n}", "answer_rows": _answers(*who)}
            for n, who in enumerate(per_activity, start=1)]


def _table(per_activity: list[list[str]]) -> dict:
    acts = _acts(per_activity)
    return participation_table(acts, participation(acts, []))


def test_the_odd_activity_out_is_early() -> None:
    t = _table([["Ana"], ["Ana"], ["Ana", "Lee"], ["Lee"], ["Lee"]])
    assert (t["early_label"], t["late_label"]) == ("Early (1–3)", "Late (4–5)")
    assert {r["name"]: (r["early"], r["late"]) for r in t["rows"]} == {"Ana": (3, 0), "Lee": (1, 2)}


def test_an_even_count_splits_down_the_middle_and_a_late_only_person_has_no_early() -> None:
    t = _table([["Ana", "Sam"], ["Ana"], ["Sam", "Lee"], ["Lee"]])
    assert (t["early_label"], t["late_label"]) == ("Early (1–2)", "Late (3–4)")
    rows = {r["name"]: r for r in t["rows"]}
    assert (rows["Lee"]["early"], rows["Lee"]["late"], rows["Lee"]["total"]) == (0, 2, 2)
    assert rows["Sam"]["cells"] == [1, 0, 1, 0]
    assert all(r["early"] + r["late"] == r["total"] for r in t["rows"])


def test_a_single_activity_is_all_early() -> None:
    t = _table([["Ana"]])
    assert (t["early_label"], t["late_label"]) == ("Early (1)", "Late")
    assert (t["rows"][0]["early"], t["rows"][0]["late"]) == (1, 0)


def test_the_pdf_appendix_opens_with_the_participation_table(isolated_env: Path) -> None:
    _, folder, run = _run(isolated_env)
    data = load_results(folder, parse_session(PLAN))
    people = participation(data["activities"], run["chat"])
    html = session_html(folder, data, people)
    assert html.index("<h2>Participation</h2>") < html.index("<h2>1 · ")
    t = participation_table(data["activities"], people)
    assert all(f"<span class='nb'>{label.partition(' ')[2]}</span>" in html for label in (t["early_label"], t["late_label"]))
    assert [r["total"] for r in t["rows"]] == [p["total"] for p in people]  # the Excel sheet prints these same rows
    assert "Participation" not in session_html(folder, {**data, "activities": [], "pages": []}, people)
    assert "Participation" not in session_html(folder, data, [])


def test_the_participants_pdf_holds_the_captures_in_run_order_and_nothing_else(isolated_env: Path) -> None:
    _, folder, run = _run(isolated_env)
    data = load_results(folder, parse_session(PLAN))
    html = activities_html(folder, data).partition("<body>")[2]  # the document, not its stylesheet
    ran = [p["item_id"] for p in data["pages"] if p["kind"] == "capture"]
    assert len(ran) == 6 and [p["item_id"] for p in data["pages"] if p["kind"] == "slide"]  # the run had slides to leave out
    assert re.findall(r"captures/([\w-]+)\.png", html) == ran
    pages = re.findall(r"<div class='page[^']*'>", html)
    assert len(pages) == 1 + len(ran) and pages[0] == "<div class='page title-page'>"  # the title page, then one page per capture
    assert data["session"]["title"] in html.split("</div>")[0] and data["session"]["date"][:10] in html.split("</div>")[0]
    assert "slides/" not in html and "appendix" not in html and "Participation" not in html
    people = participation(data["activities"], run["chat"])
    assert not any(r["name"] in html for r in people)  # no per-person name anywhere


def test_the_participants_pdf_does_not_point_to_an_appendix_it_does_not_have(isolated_env: Path) -> None:
    _, folder, _ = _run(isolated_env)
    (folder / "live" / "captures" / "act-ideas.png").unlink()
    data = load_results(folder, parse_session(PLAN))
    body = activities_html(folder, data).partition("<body>")[2]
    assert "What do we need to win?" in body and "appendix" not in body
    assert "every answer is in the appendix" in session_html(folder, data, [])


def test_the_share_pdf_holds_every_slide_and_capture_in_run_order_and_no_stats(isolated_env: Path) -> None:
    _, folder, run = _run(isolated_env)
    data = load_results(folder, parse_session(PLAN))
    html = share_html(folder, data).partition("<body>")[2]  # the document, not its stylesheet
    shown = [p for p in data["pages"] if p["kind"] in ("slide", "capture")]
    assert len(shown) == 16 and share_pages(data) == data["pages"]  # no quiz in this run: nothing to leave out
    assert re.findall(r"/(?:slides|captures)/([^/']+)'", html) == [p["file"] if p["kind"] == "slide" else f"{p['item_id']}.png" for p in shown]
    assert len(re.findall(r"<div class='page[^']*'>", html)) == len(shown) and "title-page" not in html  # no extra page
    assert "appendix" not in html and "Participation" not in html and "<table" not in html
    assert not any(r["name"] in html for r in participation(data["activities"], run["chat"]))  # no per-person name anywhere
    assert "slides/" not in activities_html(folder, data)  # the participants' copy still leaves the slides out


def test_the_share_pdf_does_not_point_to_an_appendix_it_does_not_have(isolated_env: Path) -> None:
    _, folder, _ = _run(isolated_env)
    (folder / "live" / "captures" / "act-ideas.png").unlink()
    data = load_results(folder, parse_session(PLAN))
    body = share_html(folder, data).partition("<body>")[2]
    assert "What do we need to win?" in body and "was not saved." in body and "appendix" not in body


def test_a_session_of_slides_only_still_has_a_share_pdf(isolated_env: Path) -> None:
    _, folder, _ = _run(isolated_env)
    data = load_results(folder, parse_session(PLAN))
    slides = {**data, "pages": [p for p in data["pages"] if p["kind"] == "slide"], "activities": []}
    assert len(share_pages(slides)) == 10 and activities_pages(slides) == []
    assert share_html(folder, slides).count("<div class='page") == 10


def test_sheet_titles_are_valid_and_unique() -> None:
    taken: set[str] = set()
    a = sheet_title(1, "What is our [common] enemy? / really: *", taken)
    b = sheet_title(1, "What is our [common] enemy? / really: *", taken)
    assert len(a) <= 31 and len(b) <= 31 and a != b and not set("[]:*?/\\") & set(a + b)


# ---- the session PDF (real headless Chromium) ------------------------------------

def test_session_pdf_prints_every_page_then_the_appendix(client, isolated_env: Path, caplog) -> None:
    sid, folder, _ = _run(isolated_env)
    (folder / "live" / "captures" / "act-ideas.png").unlink()  # a PNG that never rendered still gets a page
    with caplog.at_level(logging.INFO):
        r = client.post(f"/api/sessions/{sid}/exports/pdf")
    # the helper logs to stderr only; the server relays it into its own log (#37)
    assert any(m.startswith("session_pdf: ✅ session PDF written") for m in caplog.messages), caplog.messages
    assert r.status_code == 200, r.text
    body = r.json()
    assert (body["slides"], body["captures"]) == (10, 6)
    assert body["missing_images"] == ["What do we need to win?"]
    assert body["pages"] is not None and body["pages"] >= 17  # 16 pages + at least one appendix page
    assert body["file"] == "session-full.pdf"
    pdf = client.get(f"/api/sessions/{sid}/exports/session-full.pdf")
    assert pdf.status_code == 200 and pdf.content[:5] == b"%PDF-"
    assert pdf.headers["content-disposition"].startswith("inline")  # opens in a tab, never lands in Downloads
    assert client.get(f"/api/sessions/{sid}/results").json()["exports"]["pdf"]["file"] == "session-full.pdf"
    # the same click wrote the share copy: the 10 slides + 6 live results, no appendix
    assert body["share"]["file"] == "session-share.pdf" and body["share"]["pages"] == 16
    share = client.get(f"/api/sessions/{sid}/exports/session-share.pdf")
    assert share.status_code == 200 and share.content[:5] == b"%PDF-"
    assert share.headers["content-disposition"].startswith("inline")
    assert client.get(f"/api/sessions/{sid}/results").json()["exports"]["share"]["file"] == "session-share.pdf"
    assert not (folder / "exports" / "session.pdf").exists()  # the old name is not written any more
    # the same click wrote the participants' copy: a title page + the 6 live results, no slide, no appendix
    assert body["activities"]["file"] == "session-activities.pdf" and body["activities"]["pages"] == 7
    shared = client.get(f"/api/sessions/{sid}/exports/session-activities.pdf")
    assert shared.status_code == 200 and shared.content[:5] == b"%PDF-"
    assert shared.headers["content-disposition"].startswith("inline")
    assert (folder / "exports" / "session-activities.pdf").is_file()
    assert client.get(f"/api/sessions/{sid}/results").json()["exports"]["activities"]["file"] == "session-activities.pdf"


def test_an_older_export_still_opens_under_its_former_name_until_it_is_exported_again(client, isolated_env: Path) -> None:
    sid, folder, _ = _run(isolated_env)
    (folder / "exports").mkdir(exist_ok=True)
    (folder / "exports" / "session.pdf").write_bytes(b"%PDF-old")
    for name in ("session.pdf", "session-full.pdf"):  # the link the Results tab builds, and the former one
        assert client.get(f"/api/sessions/{sid}/exports/{name}").content == b"%PDF-old"
    assert client.get(f"/api/sessions/{sid}/results").json()["exports"]["pdf"]["file"] == "session.pdf"
    assert client.get(f"/api/sessions/{sid}/results").json()["exports"]["share"] is None
    (folder / "exports" / "session-full.pdf").write_bytes(b"%PDF-new")  # an export since: the new name wins
    assert client.get(f"/api/sessions/{sid}/exports/session.pdf").content == b"%PDF-new"
    assert client.get(f"/api/sessions/{sid}/results").json()["exports"]["pdf"]["file"] == "session-full.pdf"


# ---- the check against Zoom's saved chat -----------------------------------------

def test_both_saved_chat_layouts_parse() -> None:
    new = "18:40:12 From Sofía L. to Everyone:\n\tSevilla, España\n\tsecond line\n18:41:00 From Carlos P. to Everyone:\n\tok\n"
    old = "18:40:12\t From  Sofía L. : Sevilla, España\n18:41:00\t From  Carlos P. : ok\n"
    a, b = parse_saved_chat(new), parse_saved_chat(old)
    assert [(m.sender, m.text) for m in a] == [("Sofía L.", "Sevilla, España\nsecond line"), ("Carlos P.", "ok")]
    assert a[0].to == "Everyone" and a[0].time == "18:40:12"
    assert [(m.sender, m.text) for m in b] == [("Sofía L.", "Sevilla, España"), ("Carlos P.", "ok")]


def test_replies_and_reactions() -> None:
    text = ('18:40:12 From Ana R. to Everyone:\n\tperfectionism\n'
            '18:40:40 From Ben S. to Everyone:\n\tReplying to "perfectionism"\n\tsame here\n'
            '18:40:50 From Ana R. to Everyone:\n\tReacted to "same here" with 👍\n')
    saved = parse_saved_chat(text)
    assert saved[1].text == "same here"
    chat = [{"sender": "Ana R.", "text": "perfectionism", "time": "18:40"}, {"sender": "Ben S.", "text": "same here", "time": "18:40"}]
    rep = reconcile(saved, chat)
    assert (rep["zoom_messages"], rep["matched"], rep["reactions"], rep["missing_count"]) == (2, 2, 1, 0)


def test_emoji_own_messages_simulated_and_time_tolerance() -> None:
    saved = parse_saved_chat("18:40:59 From Ana R. to Everyone:\n\tgreat 🎉 session\n"
                             "18:41:10 From Demo Host to Everyone:\n\tThanks!\n"
                             "18:42:00 From Ben S. to Everyone:\n\t🎉\n")
    chat = [
        {"sender": "Ana R.", "text": "great session", "time": "6:41 PM", "source": "zoom"},
        {"sender": "You", "text": "Thanks!", "time": "18:41", "own": True, "source": "zoom"},
        {"sender": "Ben S.", "text": "", "time": "18:42", "source": "zoom"},
        {"sender": "Robin T.", "text": "7", "time": "18:43", "source": "simulator"},
    ]
    rep = reconcile(saved, chat)
    assert (rep["matched"], rep["missing_count"], rep["extra_count"], rep["simulated"]) == (3, 0, 0, 1)


def test_missing_and_extra_are_listed() -> None:
    saved = parse_saved_chat("18:00:05 From Ana R. to Everyone:\n\tHi all\n18:40:00 From Ben S. to Everyone:\n\tmeetings\n")
    chat = [{"sender": "Ben S.", "text": "meetings", "time": "18:40"}, {"sender": "Cy D.", "text": "late one", "time": "18:59"}]
    rep = reconcile(saved, chat)
    assert rep["matched"] == 1 and rep["missing"] == [{"time": "18:00:05", "sender": "Ana R.", "to": "Everyone", "text": "Hi all"}]
    assert rep["extra"] == [{"time": "18:59", "sender": "Cy D.", "text": "late one"}]


def test_the_same_answer_twice_matches_twice_not_once() -> None:
    saved = parse_saved_chat("18:40:01 From Ana R. to Everyone:\n\tyes\n18:40:30 From Ana R. to Everyone:\n\tyes\n")
    rep = reconcile(saved, [{"sender": "Ana R.", "text": "yes", "time": "18:40"}])
    assert (rep["matched"], rep["missing_count"]) == (1, 1)


def test_clock_and_text_normalisation() -> None:
    assert minutes_of("18:40:12") == minutes_of("6:40 PM") == 18 * 60 + 40
    assert minutes_of("12:05 AM") == 5 and minutes_of("nope") is None
    assert norm_text("  Great 🎉   Session ") == "great session"


def test_reconcile_endpoint_keeps_the_report(client, isolated_env: Path) -> None:
    sid, folder, run = _run(isolated_env)
    saved = isolated_env / "meeting_saved_chat.txt"
    saved.write_text("18:00:01 From Early Bird to Everyone:\n\tHello!\n" + zoom_saved_chat(run["chat"]), encoding="utf-8")
    rep = client.post(f"/api/sessions/{sid}/reconcile", json={"path": str(saved)}).json()
    assert (rep["zoom_messages"], rep["matched"], rep["missing_count"]) == (42, 41, 1)
    assert rep["file"] == "meeting_saved_chat.txt"
    assert client.get(f"/api/sessions/{sid}/results").json()["reconciliation"]["matched"] == 41
    assert (folder / "exports" / "zoom-reconciliation.json").is_file()
    junk = isolated_env / "notes.txt"
    junk.write_text("just some notes\n", encoding="utf-8")
    assert client.post(f"/api/sessions/{sid}/reconcile", json={"path": str(junk)}).json()["error"]["code"] == "not_a_zoom_chat"
    assert client.post(f"/api/sessions/{sid}/reconcile", json={"path": str(isolated_env / "none.txt")}).status_code == 404


# ---- one person under several chat names -----------------------------------------

def _scale_capture(folder: Path, answers: list[tuple[str, str]], item_id: str = "act-split") -> None:
    """A frozen scale capture (1–5) written through the real parser, one (sender, text) per answer."""
    item = {"id": item_id, "kind": "activity", "type": "scale", "title": "Split", "question": "How are you?",
            "options": {"min": 1, "max": 5}}
    msgs = [{"id": 900 + i, "sender": s, "text": t, "time": "18:40", "received_at": 1_000 + i, "source": "zoom"}
            for i, (s, t) in enumerate(answers)]
    opts = options_with_defaults("scale", item["options"])
    frozen = {"item": item, "frozen_at": "2026-01-01T00:00:00+00:00", "windows": [[1_000, 2_000]], "names": False,
              "answers": [{**{k: m[k] for k in ("id", "sender", "text", "time", "received_at")}, "hidden": False,
                           "parsed": parser("scale").parse(m, opts)} for m in msgs],
              "result": result_for("scale", item["options"], msgs)}
    (folder / "live" / "captures" / f"{item_id}.json").write_text(json.dumps(frozen, ensure_ascii=False), encoding="utf-8")


SPLIT = [("Participant A", "2"), ("Participant B", "3"), ("Participant A.", "4"), ("iPhone de Participant A", "5")]
SPLIT_GROUP = ["iPhone de Participant A", "Participant A."]


def _split(isolated_env: Path) -> tuple[str, Path, dict]:
    sid, folder, run = _run(isolated_env)
    _scale_capture(folder, SPLIT)
    return sid, folder, run


def _activities(folder: Path) -> list[dict]:
    return load_results(folder, parse_session(PLAN))["activities"]


def _split_activity(folder: Path) -> dict:
    return next(a for a in _activities(folder) if a["id"] == "act-split")


def test_suggestions_group_a_device_name_case_accents_and_a_short_form() -> None:
    names = ["Participant A", "Participant A.", "iPhone de Participant A", "Zoë Quinn", "zoe quinn", "Ben Ortiz", "Ben O",
             "Ben", "Cy Dale", "Dee Dale"]
    groups = suggest(names, Aliases())
    assert {g["canonical"]: g["aliases"] for g in groups} == {
        "Participant A": SPLIT_GROUP,
        "Zoë Quinn": ["zoe quinn"],
        "Ben Ortiz": ["Ben", "Ben O"],
    }


def test_a_name_that_fits_two_people_is_not_suggested() -> None:
    assert suggest(["Ann", "Ann Lee", "Ann Moss", "Ann L"], Aliases()) == [{"canonical": "Ann Lee", "aliases": ["Ann L"]}]
    assert suggest(["Ann", "Annabel Lee"], Aliases()) == []
    assert suggest(["Alex", "Alexa"], Aliases()) == []
    assert suggest(["Sam Cole's iPhone", "Sam Cole", "iPad de Sam Cole"], Aliases()) == [
        {"canonical": "Sam Cole", "aliases": ["iPad de Sam Cole", "Sam Cole's iPhone"]}]


def test_a_roster_name_is_the_canonical_name_and_two_roster_names_are_never_one_person() -> None:
    assert suggest(["Ann L", "iPhone de Ann"], Aliases(), roster=["Ann Lee", "Bo Park"]) == [
        {"canonical": "Ann Lee", "aliases": ["Ann L", "iPhone de Ann"]}]
    assert suggest(["Ann L"], Aliases(), roster=["Ann Lee", "Ann Lin"]) == []
    assert suggest(["Ann Lee"], Aliases(), roster=["Ann Lee"]) == []


def test_names_already_merged_are_suggested_as_their_person() -> None:
    merged = Aliases({"Participant A": ["Participant A."]})
    assert suggest(["Participant A.", "iPhone de Participant A"], merged) == [
        {"canonical": "Participant A", "aliases": ["iPhone de Participant A"]}]


def test_merge_unmerge_and_the_file_round_trip(tmp_path: Path) -> None:
    a = Aliases()
    a.merge("Ann Lee", ["Ann L", "iPhone de Ann"])
    a.merge("Ann Lee", ["ann l"])  # case only: already merged
    a.merge("Bo Park", ["Bo"])
    a.merge("Ann Lee", ["Bo Park"])  # a person with their own aliases moves with them
    assert a.groups == {"Ann Lee": ["Ann L", "iPhone de Ann", "Bo Park", "Bo"]}
    assert a.canonical("  ann   L ") == "Ann Lee" and a.canonical("Cy") == "Cy"
    write_aliases(tmp_path, a)
    assert read_aliases(tmp_path).groups == a.groups
    a.unmerge("Bo")
    a.unmerge("Bo Park")
    assert a.groups == {"Ann Lee": ["Ann L", "iPhone de Ann"]}
    a.unmerge("Ann Lee")
    write_aliases(tmp_path, a)
    assert not a and not (tmp_path / "aliases.yaml").exists()


def test_a_damaged_aliases_file_reads_as_no_merges_but_is_never_overwritten(tmp_path: Path, caplog) -> None:
    (tmp_path / "aliases.yaml").write_text("Ann: [unclosed", encoding="utf-8")
    with caplog.at_level(logging.WARNING):
        assert not read_aliases(tmp_path)
    assert "read as no merges" in caplog.text
    with pytest.raises(AliasError) as exc:
        read_aliases(tmp_path, strict=True)
    assert exc.value.code == "aliases_unreadable"


def test_without_the_file_results_read_the_raw_names(isolated_env: Path) -> None:
    _, folder, _ = _split(isolated_env)
    a = _split_activity(folder)
    assert (a["answers"], a["people"]) == (4, 4) and a["summary"].endswith("average 3.5")
    assert {r["sender"] for r in a["answer_rows"]} == {s for s, _ in SPLIT}
    assert all(r["raw"] == "" and r["label"] == r["sender"] for r in a["answer_rows"])


def test_a_merge_counts_one_person_and_one_vote(isolated_env: Path) -> None:
    _, folder, _ = _split(isolated_env)
    write_aliases(folder, Aliases({"Participant A": SPLIT_GROUP}))
    a = _split_activity(folder)
    assert (a["answers"], a["people"]) == (4, 2)  # still 4 answers, from 2 people
    assert a["summary"].endswith("average 4.0")  # A's latest vote (5) and B's (3), not four votes
    assert {r["label"] for r in a["answer_rows"]} == {"Participant A", "Participant B", "Participant A (as Participant A.)",
                                                      "Participant A (as iPhone de Participant A)"}
    assert next(r for r in a["answer_rows"] if r["raw"] == "Participant A.")["sender"] == "Participant A"


def test_participation_sums_a_persons_answers_and_messages(isolated_env: Path) -> None:
    _, folder, run = _split(isolated_env)
    chat = [*run["chat"], {"sender": "Participant A", "text": "hi", "time": "18:40"},
            {"sender": "iPhone de Participant A", "text": "hello", "time": "18:41"}]
    raw = {p["name"]: p for p in participation(_activities(folder), chat)}
    assert raw["Participant A"]["total"] == 1 and "iPhone de Participant A" in raw
    aliases = Aliases({"Participant A": SPLIT_GROUP})
    write_aliases(folder, aliases)
    merged = {p["name"]: p for p in participation(_activities(folder), chat, aliases)}
    assert merged["Participant A"]["by_activity"]["act-split"] == 3
    assert (merged["Participant A"]["total"], merged["Participant A"]["messages"]) == (3, 2)
    assert not set(SPLIT_GROUP) & set(merged)


def test_the_results_tab_suggests_and_only_a_click_merges(client, isolated_env: Path) -> None:
    sid, folder, _ = _split(isolated_env)
    people = client.get(f"/api/sessions/{sid}/results").json()["people"]
    assert people["merges"] == {}
    assert people["suggestions"] == [{"canonical": "Participant A", "aliases": SPLIT_GROUP}]
    assert not (folder / "aliases.yaml").exists()  # a suggestion alone changes nothing

    r = client.post(f"/api/sessions/{sid}/aliases/merge", json={"canonical": "Participant A", "aliases": SPLIT_GROUP})
    assert r.status_code == 200 and r.json()["merges"] == {"Participant A": SPLIT_GROUP}
    assert yaml.safe_load((folder / "aliases.yaml").read_text(encoding="utf-8")) == {"Participant A": SPLIT_GROUP}

    data = client.get(f"/api/sessions/{sid}/results").json()
    assert data["people"]["suggestions"] == [] and data["people"]["merges"] == {"Participant A": SPLIT_GROUP}
    assert next(a for a in data["activities"] if a["id"] == "act-split")["people"] == 2

    assert client.post(f"/api/sessions/{sid}/exports/xlsx").status_code == 200
    part = [[c.value for c in row] for row in load_workbook(folder / "exports" / "report.xlsx")["Participation"].iter_rows()]
    assert [row[0] for row in part if str(row[0]).startswith(("Participant", "iPhone"))] == ["Participant A", "Participant B"]

    assert client.post(f"/api/sessions/{sid}/aliases/unmerge", json={"name": "Participant A"}).json() == {"merges": {}}
    assert not (folder / "aliases.yaml").exists()  # no file: the raw view is back
    back = client.get(f"/api/sessions/{sid}/results").json()["activities"]
    assert next(a for a in back if a["id"] == "act-split")["people"] == 4


def test_a_merge_refuses_to_overwrite_a_damaged_file(client, isolated_env: Path) -> None:
    sid, folder, _ = _split(isolated_env)
    (folder / "aliases.yaml").write_text("- not\n- a mapping\n", encoding="utf-8")
    r = client.post(f"/api/sessions/{sid}/aliases/merge", json={"canonical": "Participant A", "aliases": ["Participant A."]})
    assert r.status_code == 422 and r.json()["error"]["code"] == "aliases_unreadable"
    assert (folder / "aliases.yaml").read_text(encoding="utf-8") == "- not\n- a mapping\n"
