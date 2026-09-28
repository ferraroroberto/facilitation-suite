"""Synthetic Kahoot spreadsheet templates, laid out like Kahoot's official import template.

Rows 1–7 hold the title and instructions (no logo), row 8 the header in
columns B–H with Kahoot's exact texts (en dash included), and column A numbers
rows 9–108 from 1 to 100 even where they are empty. The questions are fake.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Optional

from openpyxl import Workbook

HEADER = [
    "Question - max 120 characters",
    "Answer 1 - max 75 characters",
    "Answer 2 - max 75 characters",
    "Answer 3 - max 75 characters",
    "Answer 4 - max 75 characters",
    "Time limit (sec) – 5, 10, 20, 30, 60, 90, 120, or 240 secs",
    "Correct answer(s) - choose at least one",
]

# question, answer 1–4, time limit, correct: an int, a comma string, two answers only, non-ASCII.
QUESTIONS: list[list[Any]] = [
    ["Which planet is known as the red planet?", "Venus", "Mars", "Jupiter", "Saturn", 20, 2],
    ["Which of these are primary colours of light?", "Red", "Yellow", "Blue", "Brown", 30, "1,3"],
    ["Is water wet?", "Yes", "No", None, None, 10, 1],
    ["¿Cuál es la capital de España?", "Madrid", "Sevilla", "Bogotá", "Lisboa", 60, "1"],
    ["Pick every even number", "2", "3", "4", "8", 240, "1, 3,4"],
]


def make_kahoot_xlsx(path: Path, rows: Optional[list[list[Any]]] = None, *, header_row: int = 8,
                     first_col: int = 2, numbers: bool = True) -> Path:
    """Write a template with ``rows`` (default ``QUESTIONS``) from the row after the header."""
    rows = QUESTIONS if rows is None else rows
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.cell(2, 2, "Quiz template")
    ws.cell(3, 2, "Add questions, at least two answer alternatives, time limit and choose correct answers (at least one).")
    ws.cell(4, 2, "Have fun creating your awesome quiz!")
    for c, text in enumerate(HEADER):
        ws.cell(header_row, first_col + c, text)
    if numbers and first_col > 1:
        for n in range(1, 101):
            ws.cell(header_row + n, first_col - 1, n)
    for r, row in enumerate(rows, start=header_row + 1):
        for c, value in enumerate(row):
            if value is not None:
                ws.cell(r, first_col + c, value)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
