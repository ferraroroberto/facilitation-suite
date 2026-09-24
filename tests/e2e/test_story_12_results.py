"""Story 12: after the session — run the demo session live end to end (the
room's answers arrive as Zoom-chat messages), then open Results: every
activity in the order it happened with its frozen visual, the answers with
names, the check against Zoom's saved chat, the session PDF and the Excel
report."""

from __future__ import annotations

import time
import uuid

from openpyxl import load_workbook
from playwright.sync_api import Page, expect

from tests.e2e.conftest import shot
from tests.fixtures.demo import RUN_ANSWERS, RUN_HIDDEN, build_demo_session, demo_person, zoom_saved_chat


def _state(page: Page, base: str) -> dict:
    return page.request.get(f"{base}/api/live").json()["state"]


def _run_live(page: Page, base: str, sid: str) -> None:
    """Every item in order; each activity captures its answers, like the room typing in Zoom."""
    assert page.request.post(f"{base}/api/live/activate", data={"session": sid}).ok
    n = len(page.request.get(f"{base}/api/live").json()["plan"]["run"]["items"])
    person = 0
    for i in range(n):
        st = _state(page, base)
        item = st["item_id"]
        answers = RUN_ANSWERS.get(item)
        if answers:
            assert page.request.post(f"{base}/api/actions/capture_toggle").ok
            rows = []
            for text in answers:
                rows.append({"sender": demo_person(person % 14), "text": text, "time": time.strftime("%H:%M")})
                person += 1
            batch = {"batch": uuid.uuid4().hex, "source": "zoom", "messages": rows}
            assert page.request.post(f"{base}/api/chat/messages", data=batch).ok
            deadline = time.time() + 15
            while time.time() < deadline and (_state(page, base).get("capture") or {}).get("answers", 0) < len(answers):
                time.sleep(0.2)
            for msg in page.request.get(f"{base}/api/chat/messages").json()["messages"]:
                if (item, msg["text"]) in RUN_HIDDEN:
                    assert page.request.post(f"{base}/api/actions/hide_message/{msg['id']}").ok
            assert page.request.post(f"{base}/api/actions/capture_toggle").ok
        if i < n - 1:
            assert page.request.post(f"{base}/api/actions/next").ok


def test_results_after_the_session(page: Page, webapp, shots, tmp_path) -> None:
    base = webapp.base_url
    folder = webapp.root / "sessions" / "demo" / "results-story"
    sid, _ = build_demo_session(folder, webapp.root / "sessions.local.yaml")
    _run_live(page, base, sid)

    pngs = [folder / "live" / "captures" / f"{k}.png" for k in RUN_ANSWERS]
    deadline = time.time() + 90
    while not all(p.is_file() for p in pngs) and time.time() < deadline:
        time.sleep(0.5)
    assert all(p.is_file() for p in pngs), [p.name for p in pngs if not p.is_file()]

    # Zoom saved the whole chat at meeting end (synthetic): check the app's record against it
    chat = page.request.get(f"{base}/api/chat/messages").json()["messages"]
    saved = tmp_path / "meeting_saved_chat.txt"
    saved.write_text(zoom_saved_chat(chat), encoding="utf-8")
    rep = page.request.post(f"{base}/api/sessions/{sid}/reconcile", data={"path": str(saved)}).json()
    assert (rep["zoom_messages"], rep["matched"], rep["missing_count"]) == (40, 40, 0)
    assert page.request.post(f"{base}/api/live/deactivate").ok

    page.set_viewport_size({"width": 1440, "height": 900})
    page.add_init_script(f"localStorage.setItem('facilitation-suite.session', '{sid}')")
    page.goto(base + "/")
    page.click("#tabResults")
    pane = page.locator("#paneResults")
    expect(pane.locator(".check-banner")).to_contain_text("40 of 40 messages matched")
    expect(pane.locator(".result-row")).to_have_count(6)
    expect(pane.locator(".result-row").first).to_contain_text("Where are you joining from?")
    expect(pane.locator(".result-row").first).to_contain_text("1 unplaced")

    pane.locator(".result-row", has_text="kryptonite").click()
    expect(pane.locator(".result-detail h1")).to_contain_text("kryptonite")
    expect(pane.locator(".result-detail .muted").first).to_contain_text("8 answers from")
    expect(pane.locator(".result-detail .muted").first).to_contain_text("1 hidden")
    expect(pane.locator(".top-row").first).to_contain_text("meetings")
    expect(pane.locator(".hidden-answer")).to_contain_text("hello")
    img = pane.locator(".result-visual img")
    expect(img).to_be_visible()
    page.wait_for_function("img => img.complete && img.naturalWidth === 1920", arg=img.element_handle())
    expect(pane.locator(".pdf-tile.live")).to_have_count(6)
    expect(pane.locator(".pdf-head")).to_contain_text("10 slides + 6 live results")
    shot(page, shots / "story-12-results-1-desktop.png")

    # a click on a live tile in the PDF strip selects that activity
    pane.locator(".pdf-tile.live", has_text="Live scale").click()
    expect(pane.locator(".result-row.selected")).to_contain_text("inner weather")
    expect(pane.locator(".top-row").first).to_contain_text("sun")

    with page.expect_download(timeout=120000) as dl:
        pane.locator("[data-pdf]").click()
    assert dl.value.suggested_filename == "session.pdf"
    expect(page.locator("#toast")).to_contain_text("Session PDF ready")
    assert (folder / "exports" / "session.pdf").read_bytes()[:5] == b"%PDF-"
    expect(pane.locator(".pdf-head")).to_contain_text("pages")

    with page.expect_download() as dl:
        pane.locator("[data-xlsx]").click()
    assert dl.value.suggested_filename == "report.xlsx"
    wb = load_workbook(folder / "exports" / "report.xlsx")
    assert len(wb.sheetnames) == 8

    # a file with a message from before the reader started: the banner turns amber and lists it
    saved.write_text("17:58:01 From Early Bird to Everyone:\n\tHello from the lobby!\n" + zoom_saved_chat(chat), encoding="utf-8")
    assert page.request.post(f"{base}/api/sessions/{sid}/reconcile", data={"path": str(saved)}).ok
    page.reload()
    expect(pane.locator(".check-banner.warn")).to_contain_text("40 of 41 messages matched · 1 missing")
    pane.locator("details.card--collapsible summary").click()
    expect(pane.locator(".diff-row")).to_contain_text("Hello from the lobby!")

    page.set_viewport_size({"width": 390, "height": 844})
    page.reload()
    expect(pane.locator(".result-row")).to_have_count(6)
    shot(page, shots / "story-12-results-2-phone.png")
