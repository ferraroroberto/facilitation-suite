"""Quiz results (#55): the Results tab, the Excel report and the session PDF, from ``live/quiz.jsonl``."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from src.jsonl import read_jsonl
from src.live.plan import build_run
from src.quiz.engine import CHAT, PHONE, Game, points
from src.quiz.results import quiz_results
from src.quiz.service import replay, scopes_of
from src.results.collect import load_results, participation, timeline
from src.results.excel import build_report, quiz_sheet_title
from src.results.pdf import session_html
from src.sessions.model import parse_session
from tests.fixtures.demo import PLAN, build_demo_session, write_demo_run
from tests.fixtures.quiz_run import (
    KICKED,
    PLAN_WITH_QUIZ,
    QUIZ_JSONL,
    add_quiz_section,
    generate,
    write_quiz_run,
)


def _quiz_run(isolated_env: Path, *, demo: bool = True) -> tuple[str, Path, dict[str, Any]]:
    sid, folder = build_demo_session(isolated_env / "sessions" / "demo" / "quiz", isolated_env / "sessions.local.yaml")
    if demo:
        write_demo_run(folder)
    write_quiz_run(folder)
    return sid, folder, load_results(folder, parse_session(PLAN_WITH_QUIZ))


def _board(q: dict[str, Any]) -> list[tuple]:
    return [(r["rank"], r["name"], r["score"], r["correct"], r["answered"], r["avg_ms"], r["source"]) for r in q["leaderboard"]]


# ---- the fixture is the engine's own output ------------------------------------------

def test_the_committed_fixture_matches_a_fresh_game(isolated_env: Path) -> None:
    """Replaying the engine now gives the same results as the committed quiz.jsonl (ids and secrets aside)."""
    _, folder = build_demo_session(isolated_env / "sessions" / "demo" / "fresh", isolated_env / "sessions.local.yaml")
    generate(folder)
    items = build_run(parse_session(PLAN_WITH_QUIZ), None)["items"]
    fresh = quiz_results(folder, items)
    committed = folder.parent / "committed"
    (committed / "live").mkdir(parents=True)
    (committed / "live" / "quiz.jsonl").write_bytes(QUIZ_JSONL.read_bytes())
    kept = quiz_results(committed, items)
    assert [_board(q) for q in fresh] == [_board(q) for q in kept]
    assert [[x["answers"] for x in q["questions"]] for q in fresh] == [[x["answers"] for x in q["questions"]] for q in kept]


# ---- the results: one entry per game, every number from the engine ---------------------

def test_each_game_is_an_entry_labelled_by_title_and_run(isolated_env: Path) -> None:
    _, _, data = _quiz_run(isolated_env)
    one, two = data["quizzes"]
    assert (one["id"], one["label"], two["id"], two["label"]) == (
        "qz-lobby-1", "Space quiz · run 1", "qz-lobby-2", "Space quiz · run 2")
    assert (one["player_count"], one["question_count"], two["question_count"], two["planned_questions"]) == (7, 4, 2, 4)
    assert one["summary"] == "7 players · 4 questions · winner Nebula"
    assert KICKED not in {r["name"] for r in one["leaderboard"]}  # kicked: out of the results, as on the stage


def test_scores_are_the_engines_and_add_up_per_answer(isolated_env: Path) -> None:
    _, folder, data = _quiz_run(isolated_env)
    one = data["quizzes"][0]
    assert _board(one)[:4] == [
        (1, "Nebula", 3748, 4, 4, 1825, PHONE),
        (2, "Comet", 3670, 4, 4, 2500, PHONE),
        (3, "Pulsar", 2430, 2, 4, 4350, PHONE),
        (4, "Zenith", 2335, 2, 3, 6733, CHAT),
    ]
    # the engine's own standings, replayed independently of the results module
    game = replay(read_jsonl(folder / "live" / "quiz.jsonl"))["qz-lobby-1"]
    scope = scopes_of(build_run(parse_session(PLAN_WITH_QUIZ), None)["items"])["qz-lobby"]
    assert [(s.name, s.score, s.rank) for s in game.standings(scope.questions)] == [
        (r["name"], r["score"], r["rank"]) for r in one["leaderboard"]]
    # every player × every question asked, and the points add up to the score
    assert len(one["answer_rows"]) == 7 * 4
    for r in one["leaderboard"]:
        assert sum(a["points"] for a in one["answer_rows"] if a["player_id"] == r["id"]) == r["score"]
    nebula_red = next(a for a in one["answer_rows"] if a["name"] == "Nebula" and a["item_id"] == "qz-red")
    q = scope.questions["qz-red"]
    assert nebula_red == {**nebula_red, "choice": 2, "choice_text": "Mars", "correct": True, "elapsed_ms": 2400,
                          "points": points(q, True, 2400)}
    unanswered = next(a for a in one["answer_rows"] if a["name"] == "Quasar" and a["item_id"] == "qz-sun")
    assert (unanswered["choice"], unanswered["correct"], unanswered["points"]) == (None, None, 0)


def test_each_question_has_its_distribution_and_correct_answers(isolated_env: Path) -> None:
    _, _, data = _quiz_run(isolated_env)
    red, rings, _, sun = data["quizzes"][0]["questions"]
    assert [(a["n"], a["text"], a["count"], a["correct"]) for a in red["answers"]] == [
        (1, "Venus", 1, False), (2, "Mars", 4, True), (3, "Jupiter", 1, False), (4, "Saturn", 1, False)]
    assert red["answered"] == 7 and red["players"] == 7  # the kicked player's answer no longer counts
    assert [a["n"] for a in rings["answers"]] == [1, 2, 3] and rings["correct"] == [1, 3]
    assert (sun["points"], sun["answered"]) == ("none", 5)


def test_chat_players_are_marked_and_old_files_default_to_phone() -> None:
    g = Game("l-1", "l", 1)
    g.apply({"op": "join", "player_id": "a", "name": "A", "secret": "s", "at": 1})
    g.apply({"op": "join", "player_id": "b", "name": "B", "secret": "s", "at": 2, "source": CHAT})
    assert (g.players["a"].source, g.players["b"].source) == (PHONE, CHAT)


def test_a_lobby_the_stage_only_passed_is_not_a_game(isolated_env: Path) -> None:
    _, folder = build_demo_session(isolated_env / "sessions" / "demo" / "empty", isolated_env / "sessions.local.yaml")
    (folder / "live" / "quiz.jsonl").write_text(
        '{"game": "qz-lobby-1", "at": 1, "op": "game", "lobby_id": "qz-lobby", "run": 1}\n'
        '{"game": "qz-lobby-1", "at": 1, "op": "phase", "phase": "lobby", "item_id": "qz-lobby"}\n', encoding="utf-8")
    add_quiz_section(folder)
    assert load_results(folder, parse_session(PLAN_WITH_QUIZ))["quizzes"] == []


def test_a_session_without_a_quiz_has_no_quiz_results(isolated_env: Path) -> None:
    _, folder = build_demo_session(isolated_env / "sessions" / "demo" / "plain", isolated_env / "sessions.local.yaml")
    write_demo_run(folder)
    data = load_results(folder, parse_session(PLAN))
    assert (data["quizzes"], data["quiz_pages"]) == ([], 0)
    assert not any(p["kind"].startswith("quiz_") for p in data["pages"])
    html = session_html(folder, data)
    assert "page quiz" not in html and ".podium" not in html and "Quiz ·" not in html


# ---- the PDF order ----------------------------------------------------------------------

def ev(name: str, item: str | None = None, **kw: Any) -> dict:
    return {"event": name, "at": "t", **({"item_id": item} if item else {}), **kw}


def test_quiz_pages_interleave_with_slides_and_captures_where_they_happened() -> None:
    items = [{"id": "s1", "kind": "slide", "title": "One", "slide_file": "s1.png"},
             {"id": "s2", "kind": "slide", "title": "Two", "slide_file": "s2.png"}]
    events = [ev("session_live", "s1"), ev("quiz_phase", "q1", game="g-1", phase="lobby"),
              ev("quiz_phase", "q1", game="g-1", phase="question"), ev("quiz_phase", "q1", game="g-1", phase="reveal"),
              ev("item", "s2"), ev("quiz_phase", "q1", game="g-1", phase="reveal"),  # came back: still one page
              ev("capture_stop", "act"), ev("quiz_phase", "q2", game="g-1", phase="question"),
              ev("quiz_phase", "pod", game="g-1", phase="podium"), ev("quiz_phase", "pod", game="g-1", phase="podium")]
    keys = {("g-1", "q1"), ("g-1", "podium")}  # q2 was never closed: no page
    pages = timeline(events, items, {"act"}, keys)
    assert [(p["kind"], p["item_id"]) for p in pages] == [
        ("slide", "s1"), ("quiz_question", "q1"), ("slide", "s2"), ("capture", "act"), ("quiz_podium", "pod")]
    assert timeline(events, items, {"act"}) == [p for p in pages if not p["kind"].startswith("quiz_")]


def test_the_session_pdf_has_the_quiz_pages_in_order(isolated_env: Path) -> None:
    _, folder, data = _quiz_run(isolated_env)
    kinds = [(p["kind"], p.get("game"), p["item_id"]) for p in data["pages"]]
    assert kinds[-8:] == [
        ("quiz_question", "qz-lobby-1", "qz-red"), ("quiz_question", "qz-lobby-1", "qz-rings"),
        ("quiz_question", "qz-lobby-1", "qz-moons"), ("quiz_question", "qz-lobby-1", "qz-sun"),
        ("quiz_podium", "qz-lobby-1", "qz-podium"),
        ("quiz_question", "qz-lobby-2", "qz-red"), ("quiz_question", "qz-lobby-2", "qz-rings"),
        ("quiz_podium", "qz-lobby-2", "qz-podium"),
    ]
    assert kinds[-9] == ("slide", None, "slide-110") and data["quiz_pages"] == 8
    html = session_html(folder, data)
    assert html.count("<div class='page quiz'>") == 8
    order = [html.index(t) for t in ("Which planet is known as the red planet?", "Space quiz · run 1 · podium",
                                     "Space quiz · run 2 · podium", "Quiz · Space quiz · run 1")]
    assert order == sorted(order)
    assert "Zenith" in html and "chat" in html


def test_a_game_that_never_reached_its_podium_still_ends_on_its_standings(isolated_env: Path) -> None:
    _, folder = build_demo_session(isolated_env / "sessions" / "demo" / "nopodium", isolated_env / "sessions.local.yaml")
    write_quiz_run(folder)
    events = [e for e in read_jsonl(folder / "live" / "events.jsonl") if e.get("phase") != "podium"]
    (folder / "live" / "events.jsonl").write_text("".join(json.dumps(e) + "\n" for e in events), encoding="utf-8")
    pages = load_results(folder, parse_session(PLAN_WITH_QUIZ))["pages"]
    assert [(p["kind"], p["game"]) for p in pages] == [
        *[("quiz_question", "qz-lobby-1")] * 4, ("quiz_podium", "qz-lobby-1"),
        *[("quiz_question", "qz-lobby-2")] * 2, ("quiz_podium", "qz-lobby-2")]


# ---- the Excel report --------------------------------------------------------------------

def test_excel_has_two_sheets_per_game_with_the_engines_scores(isolated_env: Path) -> None:
    _, folder, data = _quiz_run(isolated_env)
    out = folder / "exports" / "report.xlsx"
    build_report(data, participation(data["activities"], read_jsonl(folder / "live" / "chat.jsonl")), out)
    wb = load_workbook(out)
    assert wb.sheetnames[-5:] == ["Quiz – Space quiz", "Quiz – Space quiz answers", "Quiz – Space quiz run 2",
                                  "Quiz – Space quiz run 2 answers", "Participation"]
    assert len(wb.sheetnames) == 12 and all(len(n) <= 31 for n in wb.sheetnames)
    rows = [[c.value for c in r] for r in wb["Quiz – Space quiz"].iter_rows()]
    head = rows.index(["Rank", "Nickname", "Score", "Correct", "Answered", "Avg response (s)", "Source"])
    board = [(r[1], r[2], r[6]) for r in rows[head + 1: head + 8]]
    assert board == [(r["name"], r["score"], r["source"]) for r in data["quizzes"][0]["leaderboard"]]
    assert ["Questions"] == rows[head + 9][:1]
    ans = [[c.value for c in r] for r in wb["Quiz – Space quiz answers"].iter_rows()]
    head = ans.index(["Nickname", "Source", "#", "Question", "Answer", "Correct", "Points", "Response (ms)"])
    assert len(ans) - head - 1 == 7 * 4
    assert sum(r[6] for r in ans[head + 1:] if r[0] == "Nebula") == 3748


def test_quiz_sheet_names_fit_and_keep_their_suffix() -> None:
    taken: set[str] = set()
    long = "A very long quiz title: about [planets] / moons?"
    a = quiz_sheet_title(long, 1, "", taken)
    b = quiz_sheet_title(long, 1, " answers", taken)
    c = quiz_sheet_title(long, 3, " answers", taken)
    d = quiz_sheet_title(long, 1, "", taken)  # the same game name twice: made unique
    for n in (a, b, c, d):
        assert len(n) <= 31 and not set("[]:*?/\\") & set(n) and n.startswith("Quiz – ")
    assert b.endswith(" answers") and c.endswith(" run 3 answers") and len({a, b, c, d}) == 4
    assert quiz_sheet_title("", 1, "", set()) == "Quiz – Quiz"


# ---- the API: a session whose only live activity was a quiz --------------------------

def test_a_quiz_only_session_exports_its_report_and_pdf(client, isolated_env: Path) -> None:
    sid, folder, _ = _quiz_run(isolated_env, demo=False)
    data = client.get(f"/api/sessions/{sid}/results").json()
    assert [q["label"] for q in data["quizzes"]] == ["Space quiz · run 1", "Space quiz · run 2"]
    assert data["activities"] == [] and data["quiz_pages"] == 8
    r = client.get(f"/api/sessions/{sid}/exports/report.xlsx")
    assert r.status_code == 200 and r.content[:2] == b"PK"
    assert load_workbook(folder / "exports" / "report.xlsx").sheetnames == [
        "Summary", "Quiz – Space quiz", "Quiz – Space quiz answers", "Quiz – Space quiz run 2",
        "Quiz – Space quiz run 2 answers", "Participation"]
    body = client.post(f"/api/sessions/{sid}/exports/pdf").json()
    assert body["quiz_pages"] == 8 and body["pages"] >= 9  # 8 quiz pages + the appendix
